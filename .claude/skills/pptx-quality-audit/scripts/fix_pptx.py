#!/usr/bin/env python3
"""
fix_pptx.py — tự sửa các lỗi bố cục phổ biến trong file .pptx bất kỳ (kể cả file không do pptxgenjs tạo).

    pip install python-pptx
    python3 fix_pptx.py deck.pptx deck-fixed.pptx [--margin 0.5] [--min-font 12] [--max-shrink 0.25]

Sửa:
  1. Phần tử lấn lề an toàn / vượt mép  -> dời vào trong lề; nếu to hơn vùng nội dung thì thu nhỏ.
  2. Chữ tràn khung (ước tính)           -> giảm cỡ chữ từng 1pt, tối đa --max-shrink (25%) và
                                            không thấp hơn --min-font, để tiêu đề vẫn là tiêu đề;
                                            nếu vẫn tràn thì nới chiều cao khung (nếu còn chỗ trong lề).
  3. Chữ nhỏ hơn --min-font              -> nâng lên --min-font.

Không đụng tới:
  - phần tử có tên chứa "deco", "bg", "bleed" (cố ý tràn mép);
  - phần tử "!!" nằm hẳn ngoài slide (trạng thái bay vào/ra của Morph);
  - phần tử bị xoay (rotation != 0) và placeholder không có tọa độ riêng.

Không bao giờ ghi đè file gốc. Chạy lại audit_pptx.py trên file kết quả để xác nhận.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from audit_pptx import EXEMPT, EMU, Para, estimate_lines  # noqa: E402

DEFAULT_PT = 18.0


def paras_of(tf) -> list[tuple]:
    """Trả về [(paragraph, Para)] với cỡ chữ lớn nhất của từng đoạn."""
    out = []
    for p in tf.paragraphs:
        sizes = [r.font.size.pt for r in p.runs if r.font.size is not None]
        size = max(sizes) if sizes else (p.font.size.pt if p.font.size else DEFAULT_PT)
        bold = any(r.font.bold for r in p.runs)
        out.append((p, Para("".join(r.text for r in p.runs), size, bool(bold), None, 1.0)))
    return out


def needed_height(paras: list[Para], avail_w: float) -> float:
    return sum(estimate_lines(p, avail_w) * p.size * 1.2 / 72 for p in paras)


def set_size(p, size: float):
    for r in p.runs:
        r.font.size = Pt(size)


def main() -> int:
    ap = argparse.ArgumentParser(description="Tự sửa lề và tràn chữ trong .pptx")
    ap.add_argument("src")
    ap.add_argument("dst")
    ap.add_argument("--margin", type=float, default=0.5)
    ap.add_argument("--min-font", type=float, default=12)
    ap.add_argument("--max-shrink", type=float, default=0.25, help="tỷ lệ giảm cỡ chữ tối đa (0.25 = 25%%)")
    a = ap.parse_args()
    if Path(a.src).resolve() == Path(a.dst).resolve():
        sys.exit("Không ghi đè file gốc: chọn tên file kết quả khác.")

    prs = Presentation(a.src)
    W, H = prs.slide_width / EMU, prs.slide_height / EMU
    M = a.margin
    log = []

    for n, slide in enumerate(prs.slides, start=1):
        for sh in slide.shapes:
            if sh.left is None or sh.width is None or EXEMPT.search(sh.name or ""):
                continue
            if getattr(sh, "rotation", 0):
                continue
            x, y, w, h = sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU
            off_slide = x + w <= 0 or y + h <= 0 or x >= W or y >= H
            if off_slide and sh.name.startswith("!!"):
                continue
            if w < 0.05 or h < 0.05:
                continue  # trạng thái "ẩn" (vd. thẻ lật bị bóp)

            # 1. Lề
            nw, nh = min(w, W - 2 * M), min(h, H - 2 * M)
            nx = min(max(x, M), W - M - nw)
            ny = min(max(y, M), H - M - nh)
            if (round(nx, 3), round(ny, 3), round(nw, 3), round(nh, 3)) != (round(x, 3), round(y, 3), round(w, 3), round(h, 3)):
                sh.left, sh.top, sh.width, sh.height = Emu(int(nx * EMU)), Emu(int(ny * EMU)), Emu(int(nw * EMU)), Emu(int(nh * EMU))
                log.append(f'slide {n} "{sh.name}": dời/thu vào lề ({x:.2f},{y:.2f},{w:.2f}x{h:.2f}) -> '
                           f'({nx:.2f},{ny:.2f},{nw:.2f}x{nh:.2f})')
                x, y, w, h = nx, ny, nw, nh

            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue
            tf = sh.text_frame
            li = (tf.margin_left if tf.margin_left is not None else 91440) / EMU
            ri = (tf.margin_right if tf.margin_right is not None else 91440) / EMU
            ti = (tf.margin_top if tf.margin_top is not None else 45720) / EMU
            bi = (tf.margin_bottom if tf.margin_bottom is not None else 45720) / EMU
            avail_w = max(w - li - ri, 0.01)

            # 3. Cỡ chữ tối thiểu
            pairs = paras_of(tf)
            for p, pa in pairs:
                if pa.text.strip() and pa.size < a.min_font:
                    set_size(p, a.min_font)
                    log.append(f'slide {n} "{sh.name}": nâng cỡ chữ {pa.size:g} -> {a.min_font:g}pt')
                    pa.size = a.min_font

            # 2. Tràn chữ: giảm 1pt mỗi lần
            paras = [pa for _, pa in pairs]
            before = [pa.size for pa in paras]
            floor = [max(a.min_font, round(sz * (1 - a.max_shrink))) for sz in before]
            while needed_height(paras, avail_w) > (h - ti - bi) * 1.08 + 0.02:
                shrinkable = [pa for pa, fl in zip(paras, floor) if pa.size - 1 >= fl]
                if not shrinkable:
                    break
                for pa in shrinkable:
                    pa.size -= 1
            if [pa.size for pa in paras] != before:
                for (p, pa), old in zip(pairs, before):
                    if pa.size != old:
                        set_size(p, pa.size)
                log.append(f'slide {n} "{sh.name}": giảm cỡ chữ {max(before):g} -> {max(pa.size for pa in paras):g}pt để vừa khung')

            need = needed_height(paras, avail_w) + ti + bi
            if need > h * 1.08:
                room = H - M - y
                new_h = min(need, room)
                if new_h > h:
                    sh.height = Emu(int(new_h * EMU))
                    log.append(f'slide {n} "{sh.name}": nới chiều cao {h:.2f} -> {new_h:.2f}"')
                if need > room:
                    log.append(f'slide {n} "{sh.name}": VẪN TRÀN sau khi giảm chữ, cần rút gọn nội dung, nới rộng khung hoặc tách slide')

    prs.save(a.dst)
    print("\n".join(log) if log else "Không có gì cần sửa.")
    print(f"\nĐã lưu {a.dst}. Chạy lại audit_pptx.py để xác nhận.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
